"""
Word文档格式解析工具 - 基于XML和OOXML规范
"""
from docx import Document
from docx.oxml.ns import qn
import re
from typing import Dict, Any, List, Optional, Tuple, Union
from backend.utils.logger_config import logger

class FormatParser:
    """Word文档格式解析器，用于解析文档中的格式信息，支持主题颜色和样式继承"""
    
    def __init__(self, document: Document):
        """初始化解析器
        
        Args:
            document: python-docx Document对象
        """
        self.document = document
        self.theme_colors = self._extract_theme_colors()
        self.style_map = self._build_style_map()
    
    def _extract_theme_colors(self) -> Dict[str, str]:
        """提取主题颜色定义"""
        theme_colors = {
            # 默认的主题颜色映射
            'accent1': '4F81BD',
            'accent2': 'C0504D', 
            'accent3': '9BBB59',
            'accent4': '8064A2',
            'accent5': '4BACC6',
            'accent6': 'F79646',
            'dark1': '000000',
            'dark2': '1F497D',
            'light1': 'FFFFFF',
            'light2': 'EEECE1',
            'hyperlink': '0000FF',
            'followedHyperlink': '800080'
        }
        
        # 尝试从文档主题提取颜色
        try:
            for part in self.document.part.package.parts:
                if "theme" in part.partname.lower():
                    theme_xml = part.blob.decode('utf-8')
                    # 使用正则表达式提取颜色
                    color_matches = re.findall(r'<a:(\w+)>\s*<a:srgbClr val="([0-9A-F]{6})"/>', theme_xml)
                    for name, color in color_matches:
                        theme_colors[name] = color
                        
                    # 提取方案颜色
                    scheme_matches = re.findall(r'<a:schemeClr val="(\w+)"', theme_xml)
                    for name in scheme_matches:
                        if name not in theme_colors:
                            theme_colors[name] = '808080'  # 默认灰色
        except Exception as e:
            logger.error(f"提取主题颜色时出错: {str(e)}")
            
        return theme_colors
    
    def _build_style_map(self) -> Dict[str, Dict[str, Any]]:
        """构建样式映射，包括继承关系"""
        style_map = {}
        
        for style in self.document.styles:
            style_props = {
                'name': style.name,
                'font': None,
                'size': None,
                'color': None,
                'bold': None,
                'italic': None,
                'alignment': None,
                'line_spacing': None,
                'parent': None
            }
            
            # 获取基础样式
            if hasattr(style, 'base_style') and style.base_style:
                style_props['parent'] = style.base_style.name
            
            # 提取字体属性
            if hasattr(style, 'font'):
                style_props['font'] = style.font.name
                style_props['size'] = style.font.size.pt if style.font.size else None
                style_props['color'] = style.font.color.rgb if style.font.color and style.font.color.rgb else None
                style_props['bold'] = style.font.bold
                style_props['italic'] = style.font.italic
                
            # 提取段落格式
            if hasattr(style, 'paragraph_format'):
                pf = style.paragraph_format
                style_props['alignment'] = pf.alignment
                style_props['line_spacing'] = pf.line_spacing
            
            style_map[style.name] = style_props
            
        return style_map
    
    def resolve_theme_color(self, theme_name: str, tint: Optional[float] = None) -> Optional[str]:
        """解析主题颜色为RGB值
        
        Args:
            theme_name: 主题颜色名称
            tint: 色调调整值 (-1到1)
            
        Returns:
            RGB颜色值，格式为 "RGB(r,g,b)"
        """
        if theme_name not in self.theme_colors:
            return None
            
        hex_color = self.theme_colors[theme_name]
        r, g, b = int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
        
        # 应用色调
        if tint is not None:
            if tint < 0:
                r = int(r * (1 + tint))
                g = int(g * (1 + tint))
                b = int(b * (1 + tint))
            else:
                r = int(r + (255 - r) * tint)
                g = int(g + (255 - g) * tint)
                b = int(b + (255 - b) * tint)
                
        return f'RGB({r},{g},{b})'
    
    def get_run_format(self, run) -> Dict[str, Any]:
        """获取运行(run)的所有格式属性
        
        Args:
            run: 文档中的run对象
            
        Returns:
            包含所有格式属性的字典
        """
        format_info = {
            'text': run.text,
            'font': None,
            'size': None,
            'color': None,
            'bold': run.bold,
            'italic': run.italic,
            'underline': run.underline
        }
        
        # 从API获取格式
        if run.font:
            format_info['font'] = run.font.name
            format_info['size'] = run.font.size.pt if run.font.size else None
            format_info['color'] = run.font.color.rgb if run.font.color and run.font.color.rgb else None
            
        # 从XML获取更详细格式
        if hasattr(run, '_element') and run._element is not None:
            # 字体
            font_elems = run._element.xpath('.//w:rFonts')
            if font_elems:
                for attr in ['w:ascii', 'w:eastAsia', 'w:hAnsi']:
                    val = font_elems[0].get(qn(attr))
                    if val and not format_info['font']:
                        format_info['font'] = val
            
            # 字号
            sz_elems = run._element.xpath('.//w:sz')
            if sz_elems:
                val = sz_elems[0].get(qn('w:val'))
                if val:
                    format_info['size'] = int(val) / 2
            
            # 颜色处理
            color_elems = run._element.xpath('.//w:color')
            if color_elems:
                val = color_elems[0].get(qn('w:val'))
                if val and val != 'auto':
                    try:
                        r, g, b = int(val[0:2], 16), int(val[2:4], 16), int(val[4:6], 16)
                        format_info['color'] = f'RGB({r},{g},{b})'
                    except ValueError:
                        format_info['color'] = val
            
            # 主题颜色
            theme_color = run._element.xpath('.//w:color[@w:themeColor]')
            if theme_color:
                theme_name = theme_color[0].get(qn('w:themeColor'))
                tint_attr = theme_color[0].get(qn('w:themeShade')) or theme_color[0].get(qn('w:themeTint'))
                tint = float(tint_attr) / 255.0 if tint_attr else None
                theme_color_value = self.resolve_theme_color(theme_name, tint)
                if theme_color_value:
                    format_info['color'] = theme_color_value
        
        return format_info
    
    def get_paragraph_format(self, paragraph) -> Dict[str, Any]:
        """获取段落的所有格式属性，包括继承自样式的格式
        
        Args:
            paragraph: 文档中的段落对象
            
        Returns:
            包含所有格式属性的字典
        """
        format_info = {
            'text': paragraph.text,
            'style': paragraph.style.name if paragraph.style else 'Normal',
            'alignment': paragraph.alignment,
            'font': None,
            'size': None,
            'color': None,
            'bold': None,
            'italic': None,
            'left_indent': paragraph.paragraph_format.left_indent,
            'right_indent': paragraph.paragraph_format.right_indent,
            'first_line_indent': paragraph.paragraph_format.first_line_indent,
            'line_spacing': paragraph.paragraph_format.line_spacing,
            'runs': []
        }
        
        # 从样式解析格式
        style_chain = []
        current_style = paragraph.style
        while current_style:
            style_chain.append(current_style.name)
            if current_style.name in self.style_map:
                style_info = self.style_map[current_style.name]
                
                # 填充未设置的属性
                if format_info['font'] is None and style_info['font']:
                    format_info['font'] = style_info['font']
                if format_info['size'] is None and style_info['size']:
                    format_info['size'] = style_info['size']
                if format_info['color'] is None and style_info['color']:
                    format_info['color'] = style_info['color']
                if format_info['bold'] is None and style_info['bold'] is not None:
                    format_info['bold'] = style_info['bold']
                if format_info['italic'] is None and style_info['italic'] is not None:
                    format_info['italic'] = style_info['italic']
                if format_info['alignment'] is None and style_info['alignment'] is not None:
                    format_info['alignment'] = style_info['alignment']
                if format_info['line_spacing'] is None and style_info['line_spacing'] is not None:
                    format_info['line_spacing'] = style_info['line_spacing']
                
            # 处理样式继承
            if hasattr(current_style, 'base_style') and current_style.base_style:
                current_style = current_style.base_style
            else:
                break
        
        format_info['style_chain'] = ' -> '.join(style_chain)
        
        # 处理直接格式（run级别）
        for run in paragraph.runs:
            run_format = self.get_run_format(run)
            format_info['runs'].append(run_format)
        
        return format_info
    
    def get_table_format(self, table) -> Dict[str, Any]:
        """获取表格的格式信息
        
        Args:
            table: 文档中的表格对象
            
        Returns:
            包含表格格式信息的字典
        """
        table_info = {
            'rows': len(table.rows),
            'cols': len(table.columns),
            'style': table.style.name if hasattr(table, 'style') and table.style else 'Normal Table',
            'cells': []
        }
        
        # 分析表格单元格
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells):
                cell_info = {
                    'row': i,
                    'col': j,
                    'text': cell.text,
                    'paragraphs': []
                }
                
                # 分析单元格中的段落
                for para in cell.paragraphs:
                    if para.text.strip():
                        para_format = self.get_paragraph_format(para)
                        cell_info['paragraphs'].append(para_format)
                
                table_info['cells'].append(cell_info)
        
        return table_info
    
    def get_header_footer_format(self) -> Dict[str, Any]:
        """获取页眉页脚的格式信息
        
        Returns:
            包含页眉页脚格式信息的字典
        """
        hf_info = {
            'header': {
                'has_header': False,
                'paragraphs': []
            },
            'footer': {
                'has_footer': False,
                'paragraphs': []
            }
        }
        
        if not self.document.sections:
            return hf_info
            
        section = self.document.sections[0]
        
        # 处理页眉
        if section.header:
            header_paragraphs = section.header.paragraphs
            if any(p.text.strip() for p in header_paragraphs):
                hf_info['header']['has_header'] = True
                for para in header_paragraphs:
                    if para.text.strip():
                        para_format = self.get_paragraph_format(para)
                        hf_info['header']['paragraphs'].append(para_format)
        
        # 处理页脚
        if section.footer:
            footer_paragraphs = section.footer.paragraphs
            if any(p.text.strip() for p in footer_paragraphs):
                hf_info['footer']['has_footer'] = True
                for para in footer_paragraphs:
                    if para.text.strip():
                        para_format = self.get_paragraph_format(para)
                        hf_info['footer']['paragraphs'].append(para_format)
        
        return hf_info
    
    def analyze_document(self) -> Dict[str, Any]:
        """分析整个文档并返回格式信息
        
        Returns:
            包含文档所有格式信息的字典
        """
        doc_info = {
            'paragraphs': [],
            'tables': [],
            'header_footer': self.get_header_footer_format()
        }
        
        # 分析段落
        for para in self.document.paragraphs:
            if para.text.strip():  # 跳过空段落
                para_format = self.get_paragraph_format(para)
                doc_info['paragraphs'].append(para_format)
        
        # 分析表格
        for table in self.document.tables:
            table_info = self.get_table_format(table)
            doc_info['tables'].append(table_info)
        
        return doc_info 