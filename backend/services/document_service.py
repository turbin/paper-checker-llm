import os
import zipfile
import io
import json
from io import BytesIO
from backend.utils.logger_config import logger
from backend.utils.format_parser import FormatParser
from backend.utils.docx_format_parser import DocxFormatParser
from backend.utils.config import Config
from docx import Document
from typing import Dict, List, Any, Optional, Union, BinaryIO
from datetime import datetime

class DocumentService:
    @staticmethod
    def extract_docx_content(doc: Document, use_xml_parser: bool = False, doc_path: Optional[str] = None) -> Dict[str, Any]:
        """提取文档内容，按章节组织，使用FormatParser进行格式解析
        
        Args:
            doc: Document对象
            use_xml_parser: 是否使用基于XML的格式解析器
            doc_path: 文档路径，当use_xml_parser=True时需要提供
            
        Returns:
            按章节组织的文档内容和格式信息
        """
        content = {}
        
        try:
            # 根据参数选择解析器
            if use_xml_parser:
                if doc_path:
                    # 使用DocxFormatParser
                    parser = DocxFormatParser(doc_path)
                    # 直接使用extract_docx_content方法
                    return parser.extract_docx_content()
                else:
                    logger.warning("使用XML解析器但未提供doc_path，将降级使用标准解析器")
                    parser = FormatParser(doc)
            else:
                # 使用标准FormatParser
                parser = FormatParser(doc)
            
            # 使用FormatParser获取完整格式信息
            doc_info = parser.analyze_document()
            
            # 分析封面
            cover_content = {}
            for idx, para in enumerate(doc_info["paragraphs"]):
                if idx < 10:  # 假设前10个段落是封面
                    if "封面" in para.get("style", "") or idx < 5:
                        # 根据内容判断作用
                        text = para.get("text", "")
                        if "题目" in text or "标题" in text or para.get("size", 0) > 16:
                            cover_content["title"] = para
                        elif "作者" in text or "姓名" in text:
                            cover_content["author"] = para
                        else:
                            # 添加到其他信息
                            cover_content[f"info_{idx}"] = para
            
            content["封面"] = cover_content
            
            # 分析页眉页脚
            content["页眉页脚"] = doc_info.get("header_footer", {})
            
            # 分析表格
            content["表格"] = doc_info.get("tables", [])
            
            # 分析章节
            chapters = {}
            current_chapter = "未分类"
            current_content = []
            
            for para in doc_info["paragraphs"]:
                style = para.get("style", "")
                text = para.get("text", "")
                
                if not text.strip():
                    continue
                    
                # 检查是否是章节标题
                if "Heading" in style or "标题" in style:
                    # 保存前一章节
                    if current_content:
                        chapters[current_chapter] = current_content
                    
                    # 开始新章节
                    current_chapter = text
                    current_content = []
                elif "目录" in style:
                    # 处理目录
                    if "目录" not in content:
                        content["目录"] = []
                    content["目录"].append(para)
                else:
                    # 普通段落
                    current_content.append(para)
            
            # 保存最后一章节
            if current_content:
                chapters[current_chapter] = current_content
                
            # 合并章节内容
            for chapter_name, chapter_content in chapters.items():
                content[chapter_name] = chapter_content
                
            return content
            
        except Exception as e:
            logger.error(f"提取文档内容时出错: {str(e)}", exc_info=True)
            # 降级到基础提取方法
            return DocumentService._extract_basic_content(doc)
    
    @staticmethod
    def extract_docx_content_with_xml(doc_path: str) -> Dict[str, Any]:
        """使用基于XML/OOXML的解析器提取文档内容
        
        Args:
            doc_path: 文档路径
            
        Returns:
            按章节组织的文档内容和格式信息
        """
        try:
            parser = DocxFormatParser(doc_path)
            return parser.extract_docx_content()
        except Exception as e:
            logger.error(f"使用XML解析器提取文档内容失败: {str(e)}", exc_info=True)
            # 降级到标准解析方法
            doc = Document(doc_path)
            return DocumentService.extract_docx_content(doc)
    
    @staticmethod
    def _extract_basic_content(doc: Document) -> Dict[str, Any]:
        """基础内容提取方法（降级备用）"""
        content = {}
        current_section = "未分类"
        current_content = []
        
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
                
            # 检查是否是章节标题
            if paragraph.style.name.startswith('Heading'):
                if current_content:
                    content[current_section] = current_content
                current_section = text
                current_content = []
            else:
                current_content.append({
                    'text': text,
                    'style': paragraph.style.name,
                    'font_size': paragraph.style.font.size.pt if hasattr(paragraph.style, 'font') and paragraph.style.font and paragraph.style.font.size else None,
                    'alignment': str(paragraph.alignment)
                })
        
        # 添加最后一个章节
        if current_content:
            content[current_section] = current_content
            
        return content
    
    @staticmethod
    def read_docx(file_stream: BinaryIO) -> Document:
        """读取docx文件"""
        try:
            return Document(file_stream)
        except Exception as e:
            logger.error(f"读取docx文件失败: {str(e)}")
            raise 

    @staticmethod
    def read_zip(file_stream: BinaryIO) -> Optional[Document]:
        """读取zip文件并提取docx文件"""
        try:
            # 读取zip文件内容
            zip_data = io.BytesIO(file_stream.read())
            file_stream.seek(0)  # 重置文件指针
            
            with zipfile.ZipFile(zip_data) as zip_file:
                # 获取第一个docx文件
                docx_files = [f for f in zip_file.namelist() if f.endswith('.docx')]
                if not docx_files:
                    logger.warning('压缩包中未找到.docx文件')
                    return None
                
                logger.debug(f'找到docx文件：{docx_files[0]}')
                # 读取docx文件内容
                docx_content = io.BytesIO(zip_file.read(docx_files[0]))
                doc = Document(docx_content)
                
                return doc
        except Exception as e:
            logger.error(f"处理ZIP文件时出错: {str(e)}", exc_info=True)
            return None

    @staticmethod
    def save_temp_file(file, prefix: str = "temp_") -> str:
        """保存上传的文件到临时目录
        
        Args:
            file: 上传的文件对象
            prefix: 临时文件名前缀
            
        Returns:
            临时文件的完整路径
        """
        try:
            # 创建临时目录
            temp_dir = os.path.join(os.getcwd(), "temp")
            if not os.path.exists(temp_dir):
                os.makedirs(temp_dir)
                
            # 生成唯一文件名
            timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
            filename = f"{prefix}{timestamp}_{file.filename}"
            file_path = os.path.join(temp_dir, filename)
            
            # 保存文件
            file.save(file_path)
            logger.info(f"已保存临时文件: {file_path}")
            
            return file_path
        except Exception as e:
            logger.error(f"保存临时文件失败: {str(e)}", exc_info=True)
            raise

    def read_document(self, file, save_temp: bool = False) -> Union[Document, Dict[str, Any]]:
        """读取文档文件（支持docx和zip格式）
        
        Args:
            file: 上传的文件对象
            save_temp: 是否保存临时文件用于XML解析
            
        Returns:
            Document对象或包含Document和文件路径的字典
        """
        try:
            filename = file.filename if hasattr(file, 'filename') else 'unknown'
            file_extension = os.path.splitext(filename)[1].lower()
            
            # 保存文件内容到内存
            buffer = BytesIO()
            file.save(buffer)
            buffer.seek(0)
            
            # 如果需要XML解析，保存临时文件
            temp_file_path = None
            if save_temp:
                temp_file_path = self.save_temp_file(file)
            
            # 处理不同文件类型
            doc = None
            if file_extension == '.zip':
                doc = self.read_zip(buffer)
                if not doc:
                    raise ValueError("未在ZIP文件中找到有效的DOCX文件")
            else:
                doc = self.read_docx(buffer)
            
            # 如果保存了临时文件，返回文档和路径
            if save_temp and temp_file_path:
                return {
                    "document": doc,
                    "document_path": temp_file_path
                }
            
            return doc
            
        except Exception as e:
            logger.error(f"读取文件失败: {str(e)}")
            raise ValueError(f"不支持的文件格式: {str(e)}")
    
    @staticmethod
    def save_extracted_content(content: Dict[str, Any], output_path: str) -> None:
        """将提取的内容保存为JSON文件（用于调试）
        
        Args:
            content: 提取的文档内容
            output_path: 输出文件路径
        """
        try:
            # 处理不可序列化的对象
            def json_serializer(obj):
                if hasattr(obj, "__dict__"):
                    return obj.__dict__
                return str(obj)
            
            # 保存到文件
            with open(output_path, "w", encoding="utf-8") as f:
                json.dump(content, f, default=json_serializer, ensure_ascii=False, indent=2)
                
            logger.info(f"已保存提取内容到: {output_path}")
        except Exception as e:
            logger.error(f"保存提取内容失败: {str(e)}", exc_info=True) 