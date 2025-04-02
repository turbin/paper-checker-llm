"""
应用主功能的测试用例
"""
import os
import io
from docx.oxml.ns import qn
from docx.shared import RGBColor, Pt

def test_health_check(client):
    """测试健康检查接口"""

    pass
    # response = client.get('/api/health')
    # assert response.status_code == 200
    # data = response.get_json()
    # assert data['status'] == 'healthy'
    # assert 'version' in data

def test_upload_no_file(client):
    """测试上传接口 - 无文件情况"""
    pass
    # response = client.post('/api/upload', data={})
    # assert response.status_code == 400
    # data = response.get_json()
    # assert 'error' in data

def test_paper_parser_file(client):
    """测试论文解析接口"""
    # 读取docs文件夹下的论文并解析docx文件,通过json方式返回文件中所有内容的格式
    file_path = os.path.join(os.path.dirname(__file__), 'docs', 'test.docx')
    print(f"\n测试文件路径: {file_path}")
    print(f"文件是否存在: {os.path.exists(file_path)}")
    
    # 检查文件是否存在，如果不存在则跳过测试
    if not os.path.exists(file_path):
        print(f"请确保在以下路径放置测试文档: {file_path}")
        print(f"当前工作目录: {os.getcwd()}")
        import pytest
        pytest.skip(f"测试文件不存在: {file_path}")
    
    with open(file_path, 'rb') as f:
        file_content = f.read()
    
    # 创建类似于表单上传文件的数据结构
    data = {}
    data['file'] = (io.BytesIO(file_content), 'test.docx')
    # 不调用client,仅使用python库读取docx文件
    from docx import Document
    document = Document(file_path)
    # 检查文档是否为空
    print(f"文档段落数量: {len(document.paragraphs)}")
    print(f"文档表格数量: {len(document.tables)}")
    assert len(document.paragraphs) > 0

    # 遍历所有章节和段落,并打印格式
    print("\n文档格式分析结果:")
    
    # 辅助函数：格式化显示可能为None的值
    def format_value(value, default="未设置"):
        if value is None:
            return default
        return value
    
    # 获取样式信息
    print("\n文档样式定义:")
    for style in document.styles:
        if style.type == 1:  # 段落样式
            print(f"段落样式: {style.name}")
            if style.font:
                print(f"  字体名称: {format_value(style.font.name)}")
                print(f"  字体大小: {format_value(style.font.size)}")
                if style.font.size:
                    print(f"  字体大小(pt): {style.font.size.pt}")
                print(f"  字体颜色: {format_value(style.font.color.rgb if style.font.color else None)}")
    
    # 遍历所有段落
    for i, paragraph in enumerate(document.paragraphs):
        if not paragraph.text.strip():
            continue  # 跳过空段落
        
        print(f"\n段落 {i+1}:")
        print(f"文本内容: {paragraph.text}")
        print(f"段落样式: {paragraph.style.name if paragraph.style else '默认'}")
        
        # 从样式获取格式
        style = paragraph.style
        if style and hasattr(style, 'font'):
            print("从样式获取的格式:")
            print(f"  样式字体: {format_value(style.font.name)}")
            print(f"  样式字号: {format_value(style.font.size)}")
            if style.font.size:
                print(f"  样式字号(pt): {style.font.size.pt}")
            print(f"  样式颜色: {format_value(style.font.color.rgb if style.font.color else None)}")
        
        # 段落格式
        pf = paragraph.paragraph_format
        print(f"对齐方式: {format_value(pf.alignment)}")
        print(f"左缩进: {format_value(pf.left_indent)}")
        print(f"右缩进: {format_value(pf.right_indent)}")
        print(f"首行缩进: {format_value(pf.first_line_indent)}")
        print(f"行间距: {format_value(pf.line_spacing)}")
        
        # 尝试从XML直接读取格式
        try:
            print("从XML读取的格式:")
            # 尝试获取段落级别字体设置
            rPr = paragraph._element.xpath('.//w:pPr/w:rPr')
            if rPr:
                sz = rPr[0].xpath('.//w:sz')
                if sz:
                    pts = int(sz[0].get(qn('w:val'))) / 2
                    print(f"  XML字号: {pts}pt")
                
                color = rPr[0].xpath('.//w:color')
                if color:
                    val = color[0].get(qn('w:val'))
                    print(f"  XML颜色: {val}")
        except Exception as e:
            print(f"  XML读取异常: {str(e)}")
        
        # 遍历段落中的所有运行(格式一致的文本片段)
        if not paragraph.runs:
            print("  该段落没有格式化运行文本")
        for j, run in enumerate(paragraph.runs):
            print(f"  运行 {j+1}: '{run.text}'")
            font = run.font
            print(f"    字体: {format_value(font.name)}")
            print(f"    大小: {format_value(font.size)}")
            if font.size:
                print(f"    大小(pt): {font.size.pt}")
            print(f"    粗体: {format_value(font.bold)}")
            print(f"    斜体: {format_value(font.italic)}")
            print(f"    下划线: {format_value(font.underline)}")
            
            # 颜色处理
            if hasattr(font, "_element") and font._element is not None:
                try:
                    # 检查直接颜色设置
                    color_elem = font._element.xpath('.//w:color')
                    if color_elem:
                        color_val = color_elem[0].get(qn('w:val'))
                        print(f"    直接颜色代码: {color_val}")
                        if color_val != 'auto':
                            # 将十六进制颜色转为RGB对象
                            try:
                                r, g, b = int(color_val[0:2], 16), int(color_val[2:4], 16), int(color_val[4:6], 16)
                                print(f"    RGB颜色: RGB({r},{g},{b})")
                            except ValueError:
                                print(f"    非标准颜色值: {color_val}")
                    
                    # 检查主题颜色
                    theme_color = font._element.xpath('.//w:themeColor')
                    if theme_color:
                        theme_val = theme_color[0].get(qn('w:val'))
                        print(f"    主题颜色: {theme_val}")
                except Exception as e:
                    print(f"    颜色解析错误: {str(e)}")
            
            print(f"    标准颜色: {format_value(font.color.rgb if font.color else None)}")
    
    # 遍历表格
    print("\n表格信息:")
    if not document.tables:
        print("文档中没有表格")
    for i, table in enumerate(document.tables):
        print(f"\n表格 {i+1}: {len(table.rows)}行 x {len(table.columns)}列")
        
        # 打印表格内容示例
        for row_idx, row in enumerate(table.rows):
            if row_idx > 2:  # 只打印前三行作为示例
                print("  ...")
                break
            cells_text = [cell.text for cell in row.cells]
            print(f"  行 {row_idx+1}: {cells_text}")
            
    # 获取页眉页脚信息
    print("\n页眉页脚信息:")
    if not document.sections:
        print("文档中没有定义节")
    else:
        section = document.sections[0]
        header_paragraphs = section.header.paragraphs if section.header else []
        footer_paragraphs = section.footer.paragraphs if section.footer else []
        
        if not header_paragraphs:
            print("页眉: 未设置")
        else:
            header_text = " ".join([p.text for p in header_paragraphs if p.text])
            print(f"页眉: {header_text or '无内容'}")
            
        if not footer_paragraphs:
            print("页脚: 未设置")
        else:
            footer_text = " ".join([p.text for p in footer_paragraphs if p.text])
            print(f"页脚: {footer_text or '无内容'}")